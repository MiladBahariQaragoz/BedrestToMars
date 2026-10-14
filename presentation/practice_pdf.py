import re, sys
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor, black
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Image, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT

D='/usr/share/fonts/truetype/liberation/'
pdfmetrics.registerFont(TTFont('Sans', D+'LiberationSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('Sans-B', D+'LiberationSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('Sans-I', D+'LiberationSans-Italic.ttf'))
from reportlab.pdfbase.pdfmetrics import registerFontFamily
registerFontFamily('Sans', normal='Sans', bold='Sans-B', italic='Sans-I', boldItalic='Sans-B')

WHITE=HexColor('#ECEFF4'); GREY=HexColor('#9AA5B4'); ORANGE=HexColor('#E8743B'); AMBER=HexColor('#E5B567')
deck, out, imgdir = sys.argv[1], sys.argv[2], sys.argv[3]
p = Presentation(deck)
TITLES = {9:"Evaluation: leave-one-campaign-out validation",10:"LLM-based probabilistic classifier: input and output",
  11:"Model comparison on held-out campaigns",12:"Robustness and validity checks of the LLM prediction",
  13:"Conclusions, limitations and further work",14:"Thank you for listening"}

h1 = ParagraphStyle('h1', fontName='Sans-B', fontSize=26, leading=32, textColor=WHITE)
h2 = ParagraphStyle('h2', fontName='Sans-B', fontSize=17, leading=22, textColor=WHITE, spaceAfter=2)
meta = ParagraphStyle('meta', fontName='Sans', fontSize=12, leading=16, textColor=ORANGE, spaceAfter=8)
body = ParagraphStyle('body', fontName='Sans', fontSize=13.5, leading=19, textColor=WHITE, spaceAfter=8)
ask = ParagraphStyle('ask', parent=body, fontName='Sans-I', textColor=AMBER, leftIndent=10, borderPadding=0)
small = ParagraphStyle('small', parent=body, fontSize=13, leading=19, textColor=GREY)

def bg(c, doc):
    c.saveState(); c.setFillColor(black); c.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    c.setFont('Sans', 9); c.setFillColor(GREY)
    c.drawRightString(A4[0]-12*mm, 8*mm, f"{doc.page}")
    c.restoreState()

doc = BaseDocTemplate(out, pagesize=A4, leftMargin=12*mm, rightMargin=12*mm, topMargin=12*mm, bottomMargin=14*mm,
                      title="From Bed Rest to Mars: Milad's practice script", author="Milad Bahari Qaragoz")
doc.addPageTemplates([PageTemplate(id='bk', frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id='f')], onPage=bg)])

def esc(t): return t.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')

story = [Spacer(1, 60*mm), Paragraph("From Bed Rest to Mars", h1), Spacer(1, 4*mm),
         Paragraph("Practice script, Milad's part", h2), Spacer(1, 6*mm),
         Paragraph("Slides 9 to 14: about six minutes. Backup slides 15 to 21 are for questions only.", body),
         Paragraph("<font color='#E5B567'><i>Amber italic lines</i></font> are only for answering questions; do not read them out in the talk.", small),
         PageBreak()]

W = doc.width
for i in range(9, 22):
    s = p.slides[i-1]
    notes = s.notes_slide.notes_text_frame.text.strip()
    paras = [x.strip() for x in notes.split('\n') if x.strip()]
    timing = ""
    if paras and len(paras[0]) < 40 and paras[0].endswith('.'):
        timing = paras.pop(0).rstrip('.')
    if i >= 15:
        timing = "Backup: only if asked" + (f"  ·  {timing}" if timing else "")
    if i <= 14:
        title = TITLES[i]
    else:
        title = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().startswith('Backup'):
                title = sh.text_frame.text.strip().split('\n')[0]; break
    head = [Paragraph(f"Slide {i}  ·  {esc(title)}", h2),
            Paragraph(timing, meta)]
    img = Image(f"{imgdir}/s-{i:02d}.png", width=W, height=W*9/16)
    img_block = [img, Spacer(1, 6*mm)]
    flow = []
    for para in paras:
        # split inline "If asked:" sentences into their own amber paragraph
        parts = re.split(r'(?=If asked)', para)
        for part in parts:
            part = part.strip()
            if not part: continue
            flow.append(Paragraph(esc(part), ask if part.startswith('If asked') else body))
    story += [KeepTogether(head + img_block)] + flow + [PageBreak()]
story.pop()
doc.build(story)
